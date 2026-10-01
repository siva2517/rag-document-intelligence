"""Compare generated requirements with a business-provided ground-truth list (CSV or XLSX)."""

import csv
import io

from openpyxl import load_workbook

from app.llm import LLM

TEXT_COLUMNS = ("requirement", "requirement text", "text", "description")


def load_ground_truth(filename: str, data: bytes) -> list[str]:
    """Read requirement texts from the first sheet. The first row must be a header; the column named
    like "requirement"/"text"/"description" is used, falling back to the first column."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    else:
        raise ValueError("Ground truth must be a .csv or .xlsx file.")
    if not rows:
        return []

    header = [str(h or "").strip().lower() for h in rows[0]]
    col = next((header.index(c) for c in TEXT_COLUMNS if c in header), 0)
    return [str(r[col]).strip() for r in rows[1:] if len(r) > col and r[col] is not None and str(r[col]).strip()]


def compare(generated: list[dict], truth: list[str], llm: LLM, threshold: float) -> dict:
    """Greedy one-to-one matching on embedding cosine similarity (embeddings are unit-normalised)."""
    gen_vecs = llm.embed([g["text"] for g in generated]) if generated else []
    truth_vecs = llm.embed(truth) if truth else []
    # Highest similarity first; ties go to the earlier ground-truth row / requirement id.
    pairs = sorted(
        (-sum(a * b for a, b in zip(tv, gv)), ti, gi) for ti, tv in enumerate(truth_vecs) for gi, gv in enumerate(gen_vecs)
    )

    used_t, used_g, matches = set(), set(), []
    for neg_sim, ti, gi in pairs:
        sim = -neg_sim
        if sim < threshold:
            break
        if ti in used_t or gi in used_g:
            continue
        used_t.add(ti)
        used_g.add(gi)
        g = generated[gi]
        matches.append({"truth": truth[ti], "id": g["id"], "text": g["text"], "page": g["page"], "similarity": round(sim, 3)})

    precision = len(matches) / len(generated) if generated else 0.0
    recall = len(matches) / len(truth) if truth else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return {
        "summary": {
            "ground_truth": len(truth),
            "generated": len(generated),
            "matched": len(matches),
            "precision": round(precision, 3),
            "recall": round(recall, 3),
            "f1": round(f1, 3),
            "threshold": threshold,
        },
        "matches": matches,
        "missed": [t for i, t in enumerate(truth) if i not in used_t],
        "extra": [{"id": g["id"], "text": g["text"], "page": g["page"]} for i, g in enumerate(generated) if i not in used_g],
    }
